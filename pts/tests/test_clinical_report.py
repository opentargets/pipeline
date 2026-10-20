import datetime
import json
import subprocess
import sys
from pathlib import Path
from typing import Any

import openpyxl
import polars as pl
import pytest

from pts.transformers import clinical_report as clinical_report_module
from pts.transformers.clinical import (
    ClinicalProvider,
    ClinicalReport,
    ClinicalReportOrigin,
    ClinicalReportType,
    ClinicalSource,
    extract_ema_clinical_report,
    replace_with_llm_indications,
)
from pts.transformers.clinical.mining import read_ema_sheet
from pts.transformers.clinical.ner import CHEMICAL_MODEL, DISEASE_NER_MODEL, DRUG_NER_MODEL, DRUGTEMIST_MODEL
from pts.transformers.clinical.stop_reasons import STOP_REASONS_MODEL
from pts.transformers.clinical_report import (
    ClinicalReportFlags,
    attach_stop_reasons,
    check_snapshot_revisions,
    flag_indirect_primary_purpose,
    map_report_entities,
    order_output_columns,
    require_staged,
    staged_drug_entities,
    validate_disease,
)
from tests.clinical.mapping_fixtures import DISEASE, DRUG


def _build_reports(entries) -> ClinicalReport:
    return ClinicalReport(df=pl.DataFrame(entries))


def _disease_index() -> pl.DataFrame:
    return pl.DataFrame({
        'id': pl.Series('id', ['EFO:0009880'], dtype=pl.Utf8),
        'obsoleteTerms': pl.Series('obsoleteTerms', [[None]], dtype=pl.List(pl.Utf8)),
    })


def _report_entry(report_id: str, disease_struct: dict[str, str | None]) -> dict:
    return {
        'id': report_id,
        'phaseFromSource': 'black box warning',
        'type': ClinicalReportType.INDICATION.value,
        'origin': ClinicalReportOrigin.DRUG_LABEL.value,
        'source': ClinicalSource.DailyMed.value,
        'provider': ClinicalProvider.CHEMBL.value,
        'year': None,
        'countries': ['United States'],
        'url': 'https://example.org',
        'drugs': [{'drugFromSource': 'BENAZEPRIL', 'drugId': 'CHEMBL1694'}],
        'diseases': [disease_struct],
    }


def test_validate_disease_null_if_only_null_structs() -> None:
    """Test that validate_disease returns None when all disease structs are null."""
    reports = _build_reports([
        _report_entry('null_disease', {'diseaseFromSource': None, 'diseaseId': None}),
        _report_entry(
            'with_disease',
            {'diseaseFromSource': 'teratogenicity', 'diseaseId': 'EFO:0009880'},
        ),
    ])

    validated = validate_disease(reports, disease_index=_disease_index())
    diseases = validated.df.filter(pl.col('id') == 'null_disease').select('diseases').to_series().to_list()[0]

    assert diseases is None


def test_validate_disease_preserves_populated_diseases() -> None:
    """Test that validate_disease preserves populated disease structs."""
    reports = _build_reports([
        _report_entry(
            'with_disease',
            {'diseaseFromSource': 'teratogenicity', 'diseaseId': 'EFO:0009880'},
        ),
    ])

    validated = validate_disease(reports, disease_index=_disease_index())
    diseases = validated.df.filter(pl.col('id') == 'with_disease').select('diseases').to_series().to_list()[0]

    assert diseases == [
        {'diseaseFromSource': 'teratogenicity', 'diseaseId': 'EFO:0009880'},
    ]


def _report(id_: str, source: str, primary_purpose: str | None = None) -> dict:
    return {
        'id': id_,
        'phaseFromSource': 'phase 3',
        'type': ClinicalReportType.INDICATION.value,
        'source': source,
        'provider': (
            ClinicalProvider.AACT.value
            if source == ClinicalSource.CLINICAL_TRIALS_GOV.value
            else ClinicalProvider.CHEMBL.value
        ),
        'origin': (
            ClinicalReportOrigin.CLINICAL_TRIAL.value
            if source == ClinicalSource.CLINICAL_TRIALS_GOV.value
            else ClinicalReportOrigin.DRUG_LABEL.value
        ),
        'trialPrimaryPurpose': primary_purpose,
        'drugs': [{'drugFromSource': 'BENAZEPRIL', 'drugId': 'CHEMBL1694'}],
        'diseases': [{'diseaseFromSource': 'hypertension', 'diseaseId': 'EFO:0000537'}],
    }


def _flagged(reports: ClinicalReport, report_id: str) -> bool:
    qc = reports.df.filter(pl.col('id') == report_id).select('qualityControls').to_series().to_list()[0]
    return qc is not None and ClinicalReportFlags.INDIRECT_PRIMARY_PURPOSE.value in qc


def test_flag_indirect_primary_purpose_device_feasibility() -> None:
    reports = _build_reports([
        _report('r1', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose='TREATMENT'),
        _report('r2', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose='DEVICE_FEASIBILITY'),
        _report('r3', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose='DIAGNOSTIC'),
        _report('r4', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose='OTHER'),
    ])
    result = flag_indirect_primary_purpose(reports)
    assert not _flagged(result, 'r1')
    assert _flagged(result, 'r2')
    assert _flagged(result, 'r3')
    assert not _flagged(result, 'r4')


def test_flag_indirect_primary_purpose_no_primary_purpose() -> None:
    """Reports without a trialPrimaryPurpose should not be flagged."""
    reports = _build_reports([
        _report('r1', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose=None),
    ])
    result = flag_indirect_primary_purpose(reports)
    assert not _flagged(result, 'r1')


def test_flag_indirect_primary_purpose_with_llm_no_match_is_flagged() -> None:
    """When llm_batch_results is provided but report has no match, drug_intent is null → flagged."""
    reports = _build_reports([
        _report('r1', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose='TREATMENT'),
    ])
    llm_batch_results = pl.DataFrame({
        'id': ['r_other'],
        'drug_intent': ['therapeutic'],
    })
    result = flag_indirect_primary_purpose(reports, llm_drug_intent=llm_batch_results)
    assert _flagged(result, 'r1')


def test_flag_indirect_primary_purpose_with_llm_therapeutic_not_flagged() -> None:
    """drug_intent='therapeutic' should not be flagged."""
    reports = _build_reports([
        _report('r1', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose='TREATMENT'),
    ])
    llm_batch_results = pl.DataFrame({'id': ['r1'], 'drug_intent': ['therapeutic']})
    result = flag_indirect_primary_purpose(reports, llm_drug_intent=llm_batch_results)
    assert not _flagged(result, 'r1')


def test_flag_indirect_primary_purpose_with_llm_non_therapeutic_flagged() -> None:
    """Non-therapeutic drug_intent values (prevention, supportive_care, etc.) should be flagged."""
    reports = _build_reports([
        _report('r1', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose='TREATMENT'),
        _report('r2', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose='TREATMENT'),
    ])
    llm_batch_results = pl.DataFrame({
        'id': ['r1', 'r2'],
        'drug_intent': ['therapeutic', 'prevention'],
    })
    result = flag_indirect_primary_purpose(reports, llm_drug_intent=llm_batch_results)
    assert not _flagged(result, 'r1')
    assert _flagged(result, 'r2')


def test_flag_indirect_primary_purpose_drops_drug_intent() -> None:
    """The drug_intent column must be dropped from the output."""
    reports = _build_reports([
        _report('r1', ClinicalSource.CLINICAL_TRIALS_GOV.value, primary_purpose='TREATMENT'),
    ])
    llm_batch_results = pl.DataFrame({'id': ['r1'], 'drug_intent': ['therapeutic']})
    result = flag_indirect_primary_purpose(reports, llm_drug_intent=llm_batch_results)
    assert 'drug_intent' not in result.df.columns


def test_flag_indirect_primary_purpose_non_aact_not_flagged() -> None:
    """Non-AACT sources should not be flagged based on drug_intent, even if null or non-therapeutic."""
    reports = _build_reports([
        _report('r1', ClinicalSource.DailyMed.value, primary_purpose=None),
        _report('r2', ClinicalSource.DailyMed.value, primary_purpose=None),
        _report('r3', ClinicalSource.DailyMed.value, primary_purpose=None),
    ])
    llm_batch_results = pl.DataFrame({
        'id': ['r1', 'r2'],
        'drug_intent': ['therapeutic', 'prevention'],
    })
    result = flag_indirect_primary_purpose(reports, llm_drug_intent=llm_batch_results)
    assert not _flagged(result, 'r1')
    assert not _flagged(result, 'r2')
    assert not _flagged(result, 'r3')


def _extraction(nct_id: str, disease: str, drug: str) -> pl.DataFrame:
    """A minimal frame in `ClinicalReportExtractionSchema` shape, as `parse_batch_results` returns."""
    return pl.DataFrame({
        'id': [nct_id],
        'drug_intent': ['therapeutic'],
        'drug_intent_confidence': [0.99],
        'primary_indications': [[{'name': disease, 'evidence_quote': 'q'}]],
        'investigated_drugs': [[{'drug': drug, 'evidence_quote': 'q'}]],
    })


def test_llm_extractions_are_passed_in_their_raw_schema_shape() -> None:
    """`clinical_report` must hand the extraction to the AACT provider unprojected."""
    studies = pl.DataFrame({
        'nct_id': ['NCT0001'],
        'diseaseFromSource': ['from-aact'],
        'drugFromSource': ['from-aact'],
    })

    result = replace_with_llm_indications(studies, _extraction('NCT0001', 'asthma', 'aspirin'))

    assert result['diseaseFromSource'].to_list() == ['asthma']
    assert result['drugFromSource'].to_list() == ['aspirin']

    pre_projected = pl.DataFrame({'id': ['NCT0001'], 'drugs': [['aspirin']], 'diseases': [['asthma']]})
    with pytest.raises(pl.exceptions.ColumnNotFoundError, match='primary_indications'):
        replace_with_llm_indications(studies, pre_projected)


EMA_BANNER = [
    'Content type:',
    'Medicine',
    'Output automatically generated from content on www.ema.europa.eu on:',
    '27/05/2026 - 06:15',
]
EMA_HEADER = [
    'Category',
    'Name of medicine',
    'EMA product number',
    'Medicine status',
    'International non-proprietary name (INN) / common name',
    'Active substance',
    'Therapeutic area (MeSH)',
    'Therapeutic indication',
    'Marketing authorisation date',
    'Medicine URL',
    'Revision number',
]
EMA_ROWS = [
    [
        'Human',
        'Alpha',
        'EMEA/H/C/000001',
        'Authorised',
        'alphamab',
        'alphamab',
        'Asthma',
        '  treatment of asthma  ',
        '15/01/2020',
        'https://www.ema.europa.eu/en/medicines/human/EPAR/alpha',
        3,
    ],
    [
        'Human',
        'Beta',
        'EMEA/H/C/000002',
        'Withdrawn',
        'N/A',
        'betanib',
        None,
        'treatment of copd',
        datetime.datetime(2019, 3, 4),
        'https://www.ema.europa.eu/en/medicines/human/EPAR/beta',
        2.5,
    ],
    [
        'Human',
        'Gamma',
        'EMEA/H/C/000003',
        'Authorised',
        'gammastat',
        'gammastat',
        None,
        'nothing to extract',
        '01/02/2003',
        'https://www.ema.europa.eu/en/medicines/human/EPAR/gamma',
        None,
    ],
    [
        'Human',
        'Delta',
        'EMEA/H/C/000005',
        'Authorised',
        None,
        '   ',
        'None',
        'treatment of copd',
        '01/02/2003',
        'https://www.ema.europa.eu/en/medicines/human/EPAR/delta',
        ' \t ',
    ],
    [
        'Veterinary',
        'Vetix',
        'EMEA/V/C/000004',
        'Authorised',
        None,
        'vetixin',
        None,
        'mastitis',
        '01/02/2003',
        'https://www.ema.europa.eu/en/medicines/veterinary/EPAR/vetix',
        1,
    ],
]

EMA_NER = {
    'treatment of asthma': ['asthma'],
    'treatment of copd': ['chronic obstructive pulmonary disease'],
    'nothing to extract': [],
}

READ_EXCEL_COLUMNS = [*EMA_BANNER, *(f'__UNNAMED__{i}' for i in range(len(EMA_BANNER), len(EMA_HEADER)))]
READ_EXCEL_ROWS = [
    tuple(EMA_HEADER),
    (
        'Human',
        'Alpha',
        'EMEA/H/C/000001',
        'Authorised',
        'alphamab',
        'alphamab',
        'Asthma',
        '  treatment of asthma  ',
        '15/01/2020',
        'https://www.ema.europa.eu/en/medicines/human/EPAR/alpha',
        '3',
    ),
    (
        'Human',
        'Beta',
        'EMEA/H/C/000002',
        'Withdrawn',
        None,
        'betanib',
        None,
        'treatment of copd',
        '2019-03-04 00:00:00',
        'https://www.ema.europa.eu/en/medicines/human/EPAR/beta',
        '2.5',
    ),
    (
        'Human',
        'Gamma',
        'EMEA/H/C/000003',
        'Authorised',
        'gammastat',
        'gammastat',
        None,
        'nothing to extract',
        '01/02/2003',
        'https://www.ema.europa.eu/en/medicines/human/EPAR/gamma',
        None,
    ),
    (
        'Human',
        'Delta',
        'EMEA/H/C/000005',
        'Authorised',
        None,
        '',
        None,
        'treatment of copd',
        '01/02/2003',
        'https://www.ema.europa.eu/en/medicines/human/EPAR/delta',
        '',
    ),
    (
        'Veterinary',
        'Vetix',
        'EMEA/V/C/000004',
        'Authorised',
        None,
        'vetixin',
        None,
        'mastitis',
        '01/02/2003',
        'https://www.ema.europa.eu/en/medicines/veterinary/EPAR/vetix',
        '1',
    ),
]
"""`pl.read_excel(<bytes>, sheet_name='Medicine')` of `_ema_workbook`, captured with fastexcel 0.21.0.

Its 'N/A' and 'None' cells are set to null, as `read_ema_sheet` reads them.
"""


def _ema_workbook(path: Path, blank_rows: int = 0) -> str:
    """The EMA medicines workbook: a banner row, `blank_rows` blank rows, the header, then the medicines."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Medicine'
    for i, row in enumerate([EMA_BANNER, *[[None]] * blank_rows, EMA_HEADER, *EMA_ROWS]):
        for j, value in enumerate(row):
            sheet.cell(row=i + 1, column=j + 1, value=value)
    workbook.save(path)
    return str(path)


def _stub_ner(texts: pl.Series) -> pl.Series:
    return pl.Series([EMA_NER[text] for text in texts], dtype=pl.List(pl.String))


def _sorted_reports(reports: ClinicalReport) -> list[dict]:
    return reports.df.sort('id').with_columns(pl.col('drugs', 'diseases').list.sort()).to_dicts()


def test_ema_sheet_matches_read_excel(tmp_path: Path) -> None:
    """The sheet is `pl.read_excel`'s, except missing-value strings ('N/A', 'None', ...), null by design."""
    sheet = read_ema_sheet(_ema_workbook(tmp_path / 'ema.xlsx'))
    read_excel = pl.DataFrame(READ_EXCEL_ROWS, schema=dict.fromkeys(READ_EXCEL_COLUMNS, pl.String), orient='row')

    assert sheet.dtypes == [pl.String] * len(READ_EXCEL_COLUMNS)
    assert sheet.columns[: len(EMA_BANNER)] == EMA_BANNER
    assert sheet.rows() == READ_EXCEL_ROWS
    assert _sorted_reports(extract_ema_clinical_report(sheet, _stub_ner)) == _sorted_reports(
        extract_ema_clinical_report(read_excel, _stub_ner)
    )


def test_ema_sheet_skips_blank_rows_below_the_banner(tmp_path: Path) -> None:
    sheet = read_ema_sheet(_ema_workbook(tmp_path / 'ema.xlsx', blank_rows=7))

    assert sheet.columns == read_ema_sheet(_ema_workbook(tmp_path / 'plain.xlsx')).columns
    assert sheet.rows() == READ_EXCEL_ROWS
    assert sorted(extract_ema_clinical_report(sheet, _stub_ner).df['id'].to_list()) == [
        'emea/h/c/000001',
        'emea/h/c/000002',
        'emea/h/c/000005',
    ]


def test_ema_row_without_mesh_or_ner_dropped(tmp_path: Path) -> None:
    sheet = read_ema_sheet(_ema_workbook(tmp_path / 'ema.xlsx'))

    reports = extract_ema_clinical_report(sheet, _stub_ner).df

    assert sorted(reports['id'].to_list()) == ['emea/h/c/000001', 'emea/h/c/000002', 'emea/h/c/000005']


def test_ema_missing_value_inn_falls_back_to_the_active_substance(tmp_path: Path) -> None:
    sheet = read_ema_sheet(_ema_workbook(tmp_path / 'ema.xlsx'))

    reports = extract_ema_clinical_report(sheet, _stub_ner).df

    drugs = reports.filter(pl.col('id') == 'emea/h/c/000002')['drugs'].to_list()
    assert drugs == [[{'drugFromSource': 'betanib', 'drugId': None}]]


PUBLISHED_COLUMNS = [
    'id',
    'clinicalStage',
    'origin',
    'type',
    'source',
    'provider',
    'drugs',
    'countries',
    'trialDescription',
    'phaseFromSource',
    'trialOfficialTitle',
    'trialSponsor',
    'diseases',
    'trialLiterature',
    'sideEffects',
    'trialPhase',
    'trialWhyStopped',
    'trialStudyType',
    'qualityControls',
    'trialPrimaryPurpose',
    'title',
    'trialNumberOfArms',
    'trialOverallStatus',
    'year',
    'trialStopReasonCategories',
    'url',
    'trialStartDate',
]
"""The parquet column order of `gs://open-targets-pipeline-runs/do/platform-2610-2/output/clinical_report`."""


def test_output_column_order_matches_published() -> None:
    reports = pl.DataFrame({column: [None] for column in sorted(PUBLISHED_COLUMNS, reverse=True)})

    assert order_output_columns(reports).columns == PUBLISHED_COLUMNS


def test_output_column_order_rejects_unpublished_columns() -> None:
    reports = pl.DataFrame({column: [None] for column in [*PUBLISHED_COLUMNS, 'drug_intent']})

    with pytest.raises(ValueError, match='drug_intent'):
        order_output_columns(reports)


def test_map_report_entities_maps_each_drug_and_disease_of_a_report() -> None:
    reports = pl.DataFrame(
        [
            {
                **_report('NCT01', ClinicalSource.CLINICAL_TRIALS_GOV.value),
                'drugs': [
                    {'drugFromSource': 'Aspirin', 'drugId': None},
                    {'drugFromSource': 'Kalydeco', 'drugId': None},
                ],
                'diseases': [{'diseaseFromSource': 'Cystic fibrosis', 'diseaseId': None}],
            },
            {
                **_report('NCT02', ClinicalSource.CLINICAL_TRIALS_GOV.value),
                'drugs': [{'drugFromSource': 'Ecotrin', 'drugId': None}],
                'diseases': None,
            },
        ],
        schema_overrides={'diseases': pl.List(pl.Struct({'diseaseFromSource': pl.String, 'diseaseId': pl.String}))},
    )

    mapped = map_report_entities(
        reports,
        DISEASE,
        DRUG,
        chembl_curation=None,
        drug_entities=_staged_drug_entities(),
    )

    rows = {
        r['id']: (sorted(r['drugs'], key=str), r['diseases'])
        for r in mapped.df.select('id', 'drugs', 'diseases').iter_rows(named=True)
    }
    assert rows == {
        'nct01': (
            [
                {'drugFromSource': 'Aspirin', 'drugId': 'CHEMBL25'},
                {'drugFromSource': 'Kalydeco', 'drugId': 'CHEMBL2010601'},
            ],
            [{'diseaseFromSource': 'Cystic fibrosis', 'diseaseId': 'MONDO_0009061'}],
        ),
        'nct02': ([{'drugFromSource': 'Ecotrin', 'drugId': 'CHEMBL25'}], None),
    }


STAGED_DRUG_ENTITIES_SCHEMA = {'label': pl.String, 'extracted_drugs': pl.List(pl.String)}


def _staged_drug_entities(*rows: tuple[str, list[str]]) -> pl.DataFrame:
    return pl.DataFrame(rows, schema=STAGED_DRUG_ENTITIES_SCHEMA, orient='row')


def _staged_stop_reasons(*rows: tuple[str, list[str]]) -> pl.DataFrame:
    return pl.DataFrame(
        [(text, f'sha-{text}', categories) for text, categories in rows],
        schema={'text': pl.String, 'text_sha256': pl.String, 'categories': pl.List(pl.String)},
        orient='row',
    )


def _drug_reports(*rows: tuple[str, str, str]) -> pl.DataFrame:
    """Reports of (id, source, drug label), each with one unmapped drug and no disease."""
    return pl.DataFrame(
        [
            {
                **_report(id_, source),
                'drugs': [{'drugFromSource': label, 'drugId': None}],
                'diseases': None,
            }
            for id_, source, label in rows
        ],
        schema_overrides={'diseases': pl.List(pl.Struct({'diseaseFromSource': pl.String, 'diseaseId': pl.String}))},
    )


def _drug_ids(mapped: ClinicalReport) -> dict[tuple[str, str], list]:
    """The drug ids of each (source, drug label) of the mapped reports."""
    drugs = mapped.df.select('source', 'drugs').explode('drugs').unnest('drugs')
    return {
        (row['source'], row['drugFromSource']): sorted(row['drugId'], key=str)
        for row in drugs.group_by('source', 'drugFromSource').agg('drugId').iter_rows(named=True)
    }


def test_stop_reasons_come_from_staging() -> None:
    aact = pl.DataFrame({
        'id': ['NCT01', 'NCT02', 'NCT03', 'NCT04'],
        'trialWhyStopped': ['funding', None, '', 'funding'],
    })
    staged = _staged_stop_reasons(('funding', ['Insufficient_Funding']), ('', ['Uncategorised']), ('other', ['Other']))

    attached = attach_stop_reasons(aact, staged)

    assert attached.columns == ['id', 'trialWhyStopped', 'trialStopReasonCategories']
    assert attached.schema['trialStopReasonCategories'] == pl.List(pl.String)
    assert sorted(attached.rows()) == [
        ('NCT01', 'funding', ['Insufficient_Funding']),
        ('NCT02', None, None),
        ('NCT03', '', ['Uncategorised']),
        ('NCT04', 'funding', ['Insufficient_Funding']),
    ]


def test_missing_staged_why_stopped_fails_with_count_and_examples() -> None:
    missing = [f'reason {i}' for i in range(7)]
    aact = pl.DataFrame({'id': [f'NCT{i}' for i in range(9)], 'trialWhyStopped': [*missing, 'funding', 'funding']})
    staged = _staged_stop_reasons(('funding', ['Insufficient_Funding']))

    with pytest.raises(ValueError, match='7 stop reasons') as error:
        attach_stop_reasons(aact, staged)
    assert sum(text in str(error.value) for text in missing) == 5


def test_require_staged_ignores_nulls_and_repeats() -> None:
    require_staged(pl.Series(['a', None, 'a', '']), pl.Series(['', 'a']), 'labels')


def test_staged_names_serve_a_label_of_every_source() -> None:
    reports = _drug_reports(
        ('NCT01', ClinicalSource.CLINICAL_TRIALS_GOV.value, 'mystery pill'),
        ('TTD01', ClinicalSource.TTD.value, 'mystery pill'),
        ('X01', None, 'unknownium'),
    )

    mapped = map_report_entities(
        reports,
        DISEASE,
        DRUG,
        chembl_curation=None,
        drug_entities=_staged_drug_entities(('mystery pill', ['kalydeco']), ('unknownium', [])),
    )

    assert _drug_ids(mapped) == {
        (ClinicalSource.CLINICAL_TRIALS_GOV.value, 'mystery pill'): ['CHEMBL2010601'],
        (ClinicalSource.TTD.value, 'mystery pill'): ['CHEMBL2010601'],
        (None, 'unknownium'): [None],
    }


def test_unstaged_label_from_any_source_fails() -> None:
    reports = _drug_reports(
        ('NCT01', ClinicalSource.CLINICAL_TRIALS_GOV.value, 'mystery pill'),
        ('TTD01', ClinicalSource.TTD.value, 'ttd unstaged'),
        ('CHEMBL_IND1', ClinicalSource.DailyMed.value, 'chembl unstaged'),
        ('CHEMBL_IND2', ClinicalSource.DailyMed.value, 'Aspirin'),
    )

    with pytest.raises(ValueError, match=r"^2 drug labels not staged, e\.g\. \['chembl unstaged', 'ttd unstaged'\]$"):
        map_report_entities(
            reports,
            DISEASE,
            DRUG,
            chembl_curation=None,
            drug_entities=_staged_drug_entities(('mystery pill', ['kalydeco'])),
        )


def test_conflicting_staged_results_for_one_label_fail() -> None:
    aact = _staged_drug_entities(('pill', ['a', 'b']), ('aact only', []))
    mining = _staged_drug_entities(('pill', ['c']), ('mining only', ['m']))

    with pytest.raises(ValueError, match=r"1 drug labels staged with different results, e\.g\. \['pill'\]"):
        staged_drug_entities([aact, mining])


def test_staged_drug_entities_are_the_deduplicated_union() -> None:
    aact = _staged_drug_entities(('pill', ['b', 'a']), ('aact only', []))
    mining = _staged_drug_entities(('pill', ['a', 'b']), ('mining only', ['m']))

    staged = staged_drug_entities([aact, mining])

    assert staged.schema == pl.Schema(STAGED_DRUG_ENTITIES_SCHEMA)
    assert sorted(staged.rows()) == [('aact only', []), ('mining only', ['m']), ('pill', ['a', 'b'])]


def _metadatas(**overrides: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Metadata of the four snapshots, staged at the same revisions, with `overrides` by snapshot."""
    aact = {STOP_REASONS_MODEL: 's1', CHEMICAL_MODEL: 'c1', DRUGTEMIST_MODEL: 'd1'}
    mining = {DISEASE_NER_MODEL: 'x1', CHEMICAL_MODEL: 'c1', DRUGTEMIST_MODEL: 'd1'}
    return {
        'aact_stop_reasons': {'model': STOP_REASONS_MODEL, 'model_revision': 's1', 'model_revisions': aact},
        'aact_drug_entities': {'model': DRUG_NER_MODEL, 'model_revision': 'c1+d1', 'model_revisions': aact},
        'mining_drug_entities': {'model': DRUG_NER_MODEL, 'model_revision': 'c1+d1', 'model_revisions': mining},
        'mining_disease_entities': {'model': DISEASE_NER_MODEL, 'model_revision': 'x1', 'model_revisions': mining},
    } | overrides


def test_snapshots_with_different_revisions_fail() -> None:
    check_snapshot_revisions(_metadatas())

    drug_ner = _metadatas()['mining_drug_entities']
    with pytest.raises(ValueError, match=r'aact_drug_entities.*mining_drug_entities.*c1\+d1.*c0\+d1'):
        check_snapshot_revisions(_metadatas(mining_drug_entities=drug_ner | {'model_revision': 'c0+d1'}))

    revisions = drug_ner['model_revisions'] | {CHEMICAL_MODEL: 'c0'}
    with pytest.raises(ValueError, match=f"model_revisions.*{CHEMICAL_MODEL}.*'c1'.*'c0'"):
        check_snapshot_revisions(_metadatas(mining_drug_entities=drug_ner | {'model_revisions': revisions}))


@pytest.mark.parametrize('key', ['model', 'model_revision', 'model_revisions'])
def test_snapshot_metadata_without_revisions_is_a_clear_error(key: str) -> None:
    metadata = {k: v for k, v in _metadatas()['mining_disease_entities'].items() if k != key}

    with pytest.raises(ValueError, match=f'mining_disease_entities.*{key}'):
        check_snapshot_revisions(_metadatas(mining_disease_entities=metadata))


def test_revision_mismatch_fails_before_any_dump_is_restored(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    metadatas = _metadatas()
    metadatas['mining_drug_entities'] |= {'model_revision': 'c0+d1'}
    source = {'chembl': 'unused', 'aact': 'unused'}
    for name, metadata in metadatas.items():
        (tmp_path / name).mkdir()
        (tmp_path / name / 'metadata.json').write_text(json.dumps(metadata))
        source[name] = str(tmp_path / name)

    def no_dump(*args: object, **kwargs: object) -> None:
        raise AssertionError('dump restored')

    monkeypatch.setattr(clinical_report_module, 'read_dump_tables', no_dump)

    with pytest.raises(ValueError, match=r'c0\+d1'):
        clinical_report_module._clinical_report(source, {}, {}, None)


RELEASE_MAPPING = """
import sys

import polars as pl

from pts.transformers.clinical.mining import staged_lookup
from pts.transformers.clinical_report import map_report_entities, staged_drug_entities
from tests.clinical.mapping_fixtures import DISEASE, DRUG
from tests.test_clinical_report import _drug_reports, _staged_drug_entities

drugs = staged_drug_entities([_staged_drug_entities(('mystery pill', ['kalydeco']))])
map_report_entities(_drug_reports(('TTD01', 'TTD', 'mystery pill')), DISEASE, DRUG, None, drugs)
diseases = pl.DataFrame({'text': ['asthma'], 'extracted_diseases': [['asthma']]})
staged_lookup(diseases, 'text', 'extracted_diseases', 'disease texts')(pl.Series(['asthma']))
print(sorted({module.split('.')[0] for module in sys.modules} & {'torch', 'transformers'}))
"""


def test_release_imports_no_torch() -> None:
    result = subprocess.run(
        [sys.executable, '-c', RELEASE_MAPPING],
        cwd=Path(__file__).parents[1],
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert result.stdout.splitlines()[-1] == '[]'
