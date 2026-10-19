"""Tests for the shared dataset reader and writer."""

from __future__ import annotations

import datetime
import gzip
import json
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
import polars as pl
import pytest
from polars.testing import assert_frame_equal

from pts.transformers.utils.dataset import (
    check_destinations_absent,
    read_spreadsheet,
    scan_dataset,
    scan_datasets,
    write_dataset,
)


@pytest.fixture
def doctest_namespace():
    """Dataset I/O tests do not require the shared Spark doctest session."""
    return {}


def _write_parts(directory: Path, *frames: pl.DataFrame) -> str:
    """Write one parquet part per frame, the way an upstream step lays a dataset out."""
    directory.mkdir(parents=True, exist_ok=True)
    for index, frame in enumerate(frames):
        frame.write_parquet(directory / f'{index:08d}.parquet')
    return str(directory)


def test_read_parquet_from_a_single_file(tmp_path: Path) -> None:
    path = tmp_path / 'one.parquet'
    pl.DataFrame({'a': [1, 2]}).write_parquet(path)

    assert scan_dataset(str(path)).collect().height == 2


def test_read_parquet_from_a_directory_of_parts(tmp_path: Path) -> None:
    path = _write_parts(tmp_path / 'ds', pl.DataFrame({'a': [1]}), pl.DataFrame({'a': [2]}))

    assert sorted(scan_dataset(path).collect()['a'].to_list()) == [1, 2]


def test_read_parquet_does_not_skip_underscore_prefixed_files(tmp_path: Path) -> None:
    """Deliberate: `_`-prefixed files are NOT excluded, and this pins that.

    Spark's own markers (`_SUCCESS`, `_common_metadata`) carry no data extension, so the
    per-format glob already misses them, and we control how these datasets are produced. A
    `_`-prefixed parquet would therefore have to be put there on purpose. If someone reintroduces
    filtering, this test fails and they have to justify it rather than add it quietly.
    """
    directory = tmp_path / 'ds'
    path = _write_parts(directory, pl.DataFrame({'a': [1]}))
    pl.DataFrame({'a': [999]}).write_parquet(directory / '_hidden.parquet')

    assert sorted(scan_dataset(path).collect()['a'].to_list()) == [1, 999]


def test_read_parquet_ignores_files_of_another_format(tmp_path: Path) -> None:
    """A spark `_SUCCESS` marker, and anything else without the format's extension, is skipped."""
    directory = tmp_path / 'ds'
    path = _write_parts(directory, pl.DataFrame({'a': [1]}))
    (directory / '_SUCCESS').touch()
    (directory / 'notes.txt').write_text('ignore me')

    assert scan_dataset(path).collect()['a'].to_list() == [1]


def test_read_parquet_raises_when_a_directory_has_no_parts(tmp_path: Path) -> None:
    directory = tmp_path / 'ds'
    directory.mkdir()
    (directory / '_SUCCESS').touch()

    with pytest.raises(ValueError, match=r'no .* files found'):
        scan_dataset(str(directory))


def test_read_ndjson_from_a_gzipped_file(tmp_path: Path) -> None:
    path = tmp_path / 'x.json.gz'
    path.write_bytes(gzip.compress(json.dumps({'a': 'v'}).encode()))

    frame = scan_dataset(str(path), format='ndjson').collect()

    assert frame['a'].to_list() == ['v']


def test_read_ndjson_from_a_directory_of_parts(tmp_path: Path) -> None:
    directory = tmp_path / 'lit'
    directory.mkdir()
    for index, value in enumerate(['a', 'b']):
        (directory / f'part-{index}.json.gz').write_bytes(gzip.compress(json.dumps({'v': value}).encode()))

    frame = scan_dataset(str(directory), format='ndjson').collect()

    assert sorted(frame['v'].to_list()) == ['a', 'b']


def test_schema_pins_dtypes_and_column_order(tmp_path: Path) -> None:
    """Column order is load-bearing: it feeds de-duplication content hashes downstream."""
    path = tmp_path / 'x.json.gz'
    path.write_bytes(gzip.compress(json.dumps({'z': 1, 'a': 2}).encode()))

    schema = {'a': pl.String, 'z': pl.String}
    frame = scan_dataset(str(path), format='ndjson', schema=schema).collect()

    assert frame.columns == ['a', 'z']
    assert frame.dtypes == [pl.String, pl.String]


def test_schema_with_parquet_raises_rather_than_being_ignored(tmp_path: Path) -> None:
    """Silently ignoring a passed schema is the trap this guards.

    Polars would return the file's own dtypes with no error, so a caller pinning a schema
    would never learn it did nothing.
    """
    path = tmp_path / 'x.parquet'
    pl.DataFrame({'a': [1]}).write_parquet(path)

    with pytest.raises(ValueError, match='only applied to ndjson'):
        scan_dataset(str(path), schema={'a': pl.String})


def test_unrecognised_format_raises(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match='unrecognised format'):
        scan_dataset(str(tmp_path), format='avro')


def test_read_csv_and_tsv_differ_only_in_separator(tmp_path: Path) -> None:
    csv = tmp_path / 'x.csv'
    csv.write_text('a,b\n1,2\n')
    tsv = tmp_path / 'y.tsv'
    tsv.write_text('a\tb\n1\t2\n')

    assert scan_dataset(str(csv), format='csv').collect().to_dicts() == [{'a': 1, 'b': 2}]
    assert scan_dataset(str(tsv), format='tsv').collect().to_dicts() == [{'a': 1, 'b': 2}]


def test_read_csv_forwards_options(tmp_path: Path) -> None:
    """Delimited sources vary in ways parquet does not, so parsing options belong to the caller."""
    path = tmp_path / 'x.tsv'
    path.write_text('# a comment\n1\t2\n')

    frame = scan_dataset(
        str(path), format='tsv', has_header=False, comment_prefix='#', new_columns=['a', 'b']
    ).collect()

    assert frame.to_dicts() == [{'a': 1, 'b': 2}]


def test_read_options_rejected_for_parquet(tmp_path: Path) -> None:
    """Silently dropping them would leave a caller believing an option took effect."""
    path = tmp_path / 'x.parquet'
    pl.DataFrame({'a': [1]}).write_parquet(path)

    with pytest.raises(ValueError, match='not forwarded for parquet'):
        scan_dataset(str(path), has_header=False)


def test_read_ndjson_forwards_options(tmp_path: Path) -> None:
    """`infer_schema_length` is the reason ndjson takes options at all.

    Polars infers an ndjson schema from a bounded sample of leading rows by default, so a column
    that first appears after it is dropped SILENTLY -- no error, just a missing column. This pins
    both halves -- that the default loses the column, and that the option recovers it -- because a
    test asserting only the fix would still pass if the default were harmless.
    """
    path = tmp_path / 'x.json'
    rows = [{'a': i} for i in range(200)]
    rows.append({'a': 200, 'late': 1})
    path.write_text('\n'.join(json.dumps(row) for row in rows))

    assert 'late' not in scan_dataset(str(path), format='ndjson').collect().columns
    assert 'late' in scan_dataset(str(path), format='ndjson', infer_schema_length=None).collect().columns


def test_read_rejects_a_glob_path(tmp_path: Path) -> None:
    """A glob raises loudly rather than silently stripping to the containing directory.

    Stripping would work for `*.parquet` but silently WIDEN a selective pattern like
    `part-0*.parquet` to the whole directory -- a silent wrong-data bug, worse than a loud
    failure.
    """
    with pytest.raises(ValueError, match='looks like a glob'):
        scan_dataset(str(tmp_path / '*.parquet'))


def _codecs(path: Path) -> set[str]:
    """The compression codecs actually recorded in a parquet file's metadata."""
    import pyarrow.parquet as pq

    metadata = pq.ParquetFile(path).metadata
    return {metadata.row_group(0).column(i).compression for i in range(metadata.num_columns)}


def test_write_creates_a_directory_of_parts(tmp_path: Path) -> None:
    path = str(tmp_path / 'ds')

    write_dataset(pl.DataFrame({'a': [1, 2, 3]}), path)

    parts = sorted(Path(path).glob('*.parquet'))
    assert len(parts) == 1
    assert pl.read_parquet(parts).height == 3


def test_write_uses_zstd(tmp_path: Path) -> None:
    """Asserted from the file's own metadata, not from the argument we passed in."""
    path = str(tmp_path / 'ds')

    write_dataset(pl.DataFrame({'a': list(range(1000))}), path)

    assert _codecs(next(Path(path).glob('*.parquet'))) == {'ZSTD'}


def test_write_accepts_a_lazyframe(tmp_path: Path) -> None:
    path = str(tmp_path / 'ds')

    write_dataset(pl.LazyFrame({'a': [1, 2]}), path)

    assert pl.read_parquet(sorted(Path(path).glob('*.parquet'))).height == 2


def test_write_splits_when_over_the_threshold(tmp_path: Path) -> None:
    path = str(tmp_path / 'ds')
    frame = pl.DataFrame({'a': list(range(200_000))})

    write_dataset(frame, path, approximate_bytes_per_file=100_000)

    assert len(list(Path(path).glob('*.parquet'))) > 1


def test_write_refuses_an_occupied_destination(tmp_path: Path) -> None:
    """`PartitionBy` only ADDS numbered parts; it never clears.

    Writing into a populated directory would leave the previous run's files in place, and every
    consumer globbing `*.parquet` would read both and double-count. Refusing is chosen over
    clearing because otter has no remote delete, so a clear could only work locally and would be
    a silent no-op on the cloud destinations that matter.
    """
    directory = tmp_path / 'ds'
    directory.mkdir()
    pl.DataFrame({'a': [999]}).write_parquet(directory / 'stale.parquet')

    with pytest.raises(ValueError, match='already exists'):
        write_dataset(pl.DataFrame({'a': [1]}), str(directory))

    # the existing data is untouched -- refusing must not be destructive
    assert pl.read_parquet(sorted(directory.glob('*.parquet')))['a'].to_list() == [999]


def test_write_refuses_an_empty_but_existing_directory(tmp_path: Path) -> None:
    """Existence is the test, not emptiness: a bare directory may still be a live destination."""
    directory = tmp_path / 'ds'
    directory.mkdir()

    with pytest.raises(ValueError, match='already exists'):
        write_dataset(pl.DataFrame({'a': [1]}), str(directory))


def test_write_refuses_when_the_destination_is_a_file(tmp_path: Path) -> None:
    path = tmp_path / 'ds'
    path.write_text('not a directory')

    with pytest.raises(ValueError, match='already exists'):
        write_dataset(pl.DataFrame({'a': [1]}), str(path))


def test_round_trip(tmp_path: Path) -> None:
    path = str(tmp_path / 'ds')
    frame = pl.DataFrame({'a': [1, 2], 'b': ['x', 'y']})

    write_dataset(frame, path)

    assert scan_dataset(path).collect().sort('a').equals(frame.sort('a'))


def test_scan_datasets_reads_every_matching_dataset(tmp_path: Path) -> None:
    _write_parts(tmp_path / 'evidence_a', pl.DataFrame({'targetId': ['t1']}))
    _write_parts(tmp_path / 'evidence_b', pl.DataFrame({'targetId': ['t2']}))
    _write_parts(tmp_path / 'other', pl.DataFrame({'targetId': ['t3']}))

    frame = scan_datasets(str(tmp_path / 'evidence_*')).collect()

    assert sorted(frame['targetId'].to_list()) == ['t1', 't2']


def test_scan_datasets_unions_differing_schemas(tmp_path: Path) -> None:
    """Most evidence datasets have no `drugId`; spark's mergeSchema null-filled them."""
    _write_parts(tmp_path / 'evidence_a', pl.DataFrame({'targetId': ['t1'], 'drugId': ['d1']}))
    _write_parts(tmp_path / 'evidence_b', pl.DataFrame({'targetId': ['t2']}))

    frame = scan_datasets(str(tmp_path / 'evidence_*')).select('targetId', 'drugId').collect()

    assert sorted(zip(frame['targetId'], frame['drugId'], strict=True)) == [('t1', 'd1'), ('t2', None)]


def test_scan_datasets_raises_when_the_glob_matches_nothing(tmp_path: Path) -> None:
    """A mistyped pattern must fail loudly, not yield an empty frame.

    An empty evidence frame would silently strip every drug from the disease index rather
    than erroring, so this is the difference between a loud failure and a bad release.
    """
    with pytest.raises(ValueError, match='matched no datasets'):
        scan_datasets(str(tmp_path / 'nothing_*'))


def test_scan_datasets_requires_a_glob(tmp_path: Path) -> None:
    _write_parts(tmp_path / 'plain', pl.DataFrame({'a': [1]}))

    with pytest.raises(ValueError, match='is not a glob'):
        scan_datasets(str(tmp_path / 'plain'))


def test_scan_datasets_ignores_directories_with_no_data_files(tmp_path: Path) -> None:
    """The name glob may match a directory that carries no data at all.

    Matching is done against data FILES, not directory names (see `scan_datasets`' docstring:
    the GCS backend's glob only matches objects, so a bare directory never matches there), so an
    empty directory produces no file matches and is silently skipped rather than raising or
    contributing an empty frame.
    """
    _write_parts(tmp_path / 'evidence_a', pl.DataFrame({'targetId': ['t1']}))
    (tmp_path / 'evidence_b').mkdir()

    frame = scan_datasets(str(tmp_path / 'evidence_*')).collect()

    assert frame['targetId'].to_list() == ['t1']


def test_scan_datasets_ignores_success_markers_and_crc_sidecars(tmp_path: Path) -> None:
    """A spark `_SUCCESS` marker and a `.parquet.crc` sidecar must not become phantom datasets."""
    directory = tmp_path / 'evidence_a'
    _write_parts(directory, pl.DataFrame({'targetId': ['t1']}))
    (directory / '_SUCCESS').touch()
    (directory / '.00000000.parquet.crc').touch()

    frame = scan_datasets(str(tmp_path / 'evidence_*')).collect()

    assert frame['targetId'].to_list() == ['t1']


def test_read_spreadsheet_gives_strings_like_pandas(tmp_path: Path) -> None:
    columns = {
        'd': ['x', None, ' '],
        'a': [1, 2, 3],
        'b': [1.5, None, 2.0],
        'c': [True, False, None],
        'm': [1.0, '-', 'NaN'],
        'n': ['NA', 'n/a', ' NA'],
    }
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    ws.append(list(columns))
    for row in zip(*columns.values(), strict=True):
        ws.append(list(row))
    wb.save(tmp_path / 'cells.xlsx')

    df = read_spreadsheet(str(tmp_path / 'cells.xlsx'), 'S')

    assert df.schema == pl.Schema(dict.fromkeys(columns, pl.String))
    assert df.rows() == [
        ('x', '1', '1.5', '1', '1', None),
        (None, '2', None, '0', '-', None),
        (' ', '3', '2.0', None, None, ' NA'),
    ]


def test_spreadsheet_strings_match_old_reader(tmp_path: Path) -> None:
    """Without type inference, the strings of `pd.read_excel(dtype=str)`, as gene burden read them."""
    columns = {
        'integral': [5.0, 0.01, 0.0001, -0.0, 1e15],
        'floats': [1.5e-08, 2e-05, 0.1 + 0.2, 1 / 3, 5e-320],
        'flag': [True, False, None, True, False],
        'number_flag': [1.0, 0.0, 1, None, 1.0],
        'text': ['-', 'na', 'NA', ' a ', 'NaN'],
        'numeric_text': ['5', '5.0', 'inf', '1e-05', 'singleton'],
        'blank': [None, 'x', None, None, ''],
        'when': [datetime.datetime(2020, 1, 2), datetime.date(2021, 3, 4), None, None, None],
    }
    rows = [list(columns), *(list(row) for row in zip(*columns.values(), strict=True))]
    # a blank row between data rows, and blank rows before a formatted but empty trailing cell
    rows[3:3] = [[None] * len(columns)]
    path = _write_sheet(tmp_path / 'cells.xlsx', rows)
    trailing = openpyxl.load_workbook(path)
    trailing['S'].cell(row=len(rows) + 3, column=1).number_format = '0.00'
    trailing['S'].cell(row=len(rows) + 3, column=1).value = ''
    trailing.save(path)

    df = read_spreadsheet(path, 'S', infer_types=False)

    assert df.schema == pl.Schema(dict.fromkeys(columns, pl.String))
    assert (df.columns, df.rows()) == _old_reader(path)
    assert df.height == 6
    assert df.row(2) == (None,) * len(columns)
    assert df['integral'].to_list() == ['5', '0.01', None, '0.0001', '0', '1000000000000000']
    assert df['flag'].to_list() == ['True', 'False', None, None, 'True', 'False']


def _write_sheet(path: Path, rows: list[list[Any]], *, start: str = 'A1') -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    origin = ws[start]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            if value is not None:
                ws.cell(row=origin.row + i, column=origin.column + j, value=value)
    wb.save(path)
    return str(path)


def _cell_to_string(value: Any) -> str | None:
    """A pandas cell as gene burden's old reader passed it to spark: booleans as '1'/'0', blanks as null."""
    if pd.isna(value):
        return None
    if isinstance(value, bool | np.bool_):
        return '1' if value else '0'
    return str(value)


def _old_reader(path: str) -> tuple[list[str], list[tuple[str | None, ...]]]:
    """Gene burden's old reader: `pd.read_excel(dtype=str)` then `_cell_to_string`."""
    old = pd.read_excel(path, sheet_name='S', dtype=str).astype(object).map(_cell_to_string)
    rows = [tuple(None if pd.isna(v) else v for v in row) for row in old.itertuples(index=False)]
    return [str(c) for c in old.columns], rows


@pytest.mark.parametrize(
    ('rows', 'start'),
    [
        pytest.param([[None, None], ['a', 'b'], [1, 'x']], 'A1', id='blank_header_row'),
        pytest.param([['a', 'b'], [None, None], [1, 'x'], [2, 'y']], 'A1', id='blank_row_under_header'),
        pytest.param([['a', 'b'], [1, 'x'], [None, None], [None, None], ['note', None]], 'A1', id='footnote'),
        pytest.param([['a', 'b'], [1, 'x'], ['NA', 'nan'], [2, 'y']], 'A1', id='missing_value_row'),
        pytest.param([['a', 'b'], [1, 'x']], 'C3', id='offset_table'),
        pytest.param([['a', 'a', None, 'a.1', 'b'], [1, 2, 3, 4, None]], 'A1', id='header_names'),
        pytest.param([['a', 'b'], [1e20, '#DIV/0!'], [2**53 + 2.0, '#N/A']], 'A1', id='big_numbers_and_errors'),
        pytest.param([['a', 'b', None], ['x', None, None], ['y', 'z', None]], 'A1', id='short_rows'),
    ],
)
def test_spreadsheet_layout_matches_old_reader(tmp_path: Path, rows: list[list[Any]], start: str) -> None:
    path = _write_sheet(tmp_path / 'layout.xlsx', rows, start=start)

    df = read_spreadsheet(path, 'S', infer_types=False)

    assert (df.columns, df.rows()) == _old_reader(path)


@pytest.mark.parametrize(
    'rows',
    [
        pytest.param([['title'], ['a', None, 'b'], [None, 'x', None], [1, 2.5, 'NA'], ['note']], id='multi_row_header'),
        pytest.param([[None, None], ['a', 'b'], [None, None], [1e-8, True]], id='blank_rows'),
        pytest.param([['a', 'a', None], [5.0, '-', 'n/a']], id='repeated_names'),
    ],
)
def test_spreadsheet_without_a_header_matches_old_reader(tmp_path: Path, rows: list[list[Any]]) -> None:
    """Every row as data, as `pd.read_excel(header=None, dtype=str)` reads it."""
    path = _write_sheet(tmp_path / 'raw.xlsx', rows)

    df = read_spreadsheet(path, 'S', infer_types=False, has_header=False)

    old = pd.read_excel(path, sheet_name='S', header=None, dtype=str).astype(object).map(_cell_to_string)
    assert df.columns == [f'column_{i + 1}' for i in range(old.shape[1])]
    assert df.rows() == [tuple(None if pd.isna(v) else v for v in row) for row in old.itertuples(index=False)]


def test_spreadsheet_without_a_header_needs_text_mode(tmp_path: Path) -> None:
    path = _write_sheet(tmp_path / 'raw.xlsx', [['a'], [1]])
    with pytest.raises(ValueError, match='infer_types'):
        read_spreadsheet(path, 'S', has_header=False)


@pytest.mark.parametrize('infer_types', [True, False])
def test_read_spreadsheet_with_a_header_only(tmp_path: Path, infer_types: bool) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    ws.append(['a', 'b'])
    wb.save(tmp_path / 'header.xlsx')

    df = read_spreadsheet(str(tmp_path / 'header.xlsx'), 'S', infer_types=infer_types)

    assert df.schema == pl.Schema({'a': pl.String, 'b': pl.String})
    assert df.height == 0


def test_check_destinations_absent_names_every_existing_path(tmp_path: Path) -> None:
    (tmp_path / 'a').mkdir()
    (tmp_path / 'c').mkdir()
    paths = [str(tmp_path / name) for name in 'abc']

    with pytest.raises(ValueError, match='already exists') as error:
        check_destinations_absent(paths)
    assert str(tmp_path / 'a') in str(error.value)
    assert str(tmp_path / 'b') not in str(error.value)
    assert str(tmp_path / 'c') in str(error.value)
    check_destinations_absent([str(tmp_path / 'b')])


def test_scan_datasets_reads_only_explicit_paths(tmp_path: Path) -> None:
    """Explicit selections accept directories and single files, retaining duplicate rows."""
    first = _write_parts(tmp_path / 'selected', pl.DataFrame({'id': ['same']}))
    second = tmp_path / 'single.parquet'
    pl.DataFrame({'id': ['same']}).write_parquet(second)
    _write_parts(tmp_path / 'other', pl.DataFrame({'id': ['unselected']}))

    result = scan_datasets([first, str(second)])

    assert isinstance(result, pl.LazyFrame)
    assert_frame_equal(result.collect(), pl.DataFrame({'id': ['same', 'same']}))


def test_scan_datasets_rejects_empty_path_list() -> None:
    with pytest.raises(ValueError, match='At least one dataset path'):
        scan_datasets([])


def test_scan_datasets_explicit_paths_union_schemas_without_declared_schema(tmp_path: Path) -> None:
    first = _write_parts(tmp_path / 'a', pl.DataFrame({'id': ['a'], 'value': pl.Series([1], dtype=pl.Int32)}))
    second = _write_parts(tmp_path / 'b', pl.DataFrame({'id': ['b'], 'value': [2], 'extra': ['x']}))

    expected = pl.DataFrame({'id': ['a', 'b'], 'value': [1, 2], 'extra': [None, 'x']})
    assert_frame_equal(scan_datasets([first, second]).collect(), expected)


@pytest.mark.parametrize('selection', ['list', 'glob'])
def test_scan_datasets_matches_declared_schema(tmp_path: Path, selection: str) -> None:
    """Both overloads insert typed nulls, including missing nested struct fields."""
    first = _write_parts(tmp_path / 'dataset_a', pl.DataFrame({'id': ['a']}))
    second = _write_parts(tmp_path / 'dataset_b', pl.DataFrame({'details': [{'name': 'b'}]}))
    schema = pl.Schema({'details': pl.Struct({'name': pl.String, 'count': pl.Int32}), 'id': pl.String})
    paths = [first, second] if selection == 'list' else str(tmp_path / 'dataset_*')

    result = scan_datasets(paths, schema=schema)
    expected = pl.DataFrame({'details': [None, {'name': 'b', 'count': None}], 'id': ['a', None]}, schema=schema)

    assert isinstance(result, pl.LazyFrame)
    assert_frame_equal(result.collect(), expected)


@pytest.mark.parametrize('selection', ['list', 'glob'])
@pytest.mark.parametrize(
    'frame',
    [pl.DataFrame({'unexpected': ['x']}), pl.DataFrame({'id': [1]})],
    ids=['unexpected-column', 'mismatched-type'],
)
def test_scan_datasets_declared_schema_rejects_invalid_inputs(
    tmp_path: Path, selection: str, frame: pl.DataFrame
) -> None:
    path = _write_parts(tmp_path / 'dataset_a', frame)
    paths = [path] if selection == 'list' else str(tmp_path / 'dataset_*')

    with pytest.raises(pl.exceptions.SchemaError):
        scan_datasets(paths, schema=pl.Schema({'id': pl.String})).collect()
