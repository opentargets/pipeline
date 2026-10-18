"""Excel fixtures for the gene burden spreadsheet parsers."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl


def write_workbook(path: Path, sheets: dict[str, tuple[list[str], list[dict[str, Any]]]]) -> str:
    """Write one sheet per entry, its header first, then a row per dict; a missing key is a blank cell."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, (columns, rows) in sheets.items():
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in rows:
            sheet.append([row.get(column) for column in columns])
    workbook.save(path)
    return str(path)


def write_rows(path: Path, sheets: dict[str, list[list[Any]]]) -> str:
    """Write one sheet per entry, a row per list; None is a blank cell."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return str(path)
